"""Заполнение «Акта экспертизы внутреннего экспортного контроля» (КВЭК) и «Заключения о возможности открытого
опубликования» (ОП) по JSON с полями. Формы — 2026 г. (как в подписанном образце рег. 20260806_2, `образцы_2026/`).
Меняются только поля материала и фразы «выбрать необходимое»; подписанты, номер акта и даты остаются пустыми
(«И.О. Фамилия», подчёркивания) — их заполняют люди.

    python fill_acts.py ПОЛЯ.json ПАПКА_ВЫХОДА
    python fill_acts.py --make-op-template      # пустой шаблон ОП 2026 из подписанного образца
    python fill_acts.py --selfcheck

ПОЛЯ.json:
{
  "файл": "Доклады_ТУСУР_шумовая_модель",            # суффикс имён выходных файлов
  "вид": "статья",                                    # КВЭК: вид материала
  "вид_ОП": "статья в журнал",                        # ОП: вид материала (по списку формы)
  "название": "…",  "перевод_названия": "",          # перевод — если название на английском
  "авторы": "Моховиков Д.М., Кулинич И.В., …",       # как в публикации
  "издание": "журнал «Доклады ТУСУР»",               # именительный: ОП
  "издание_для": "журнала «Доклады ТУСУР»",          # родительный: КВЭК «предназначенный для …»
  "краткое_содержание": "…",                          # 1–3 предложения, без режимов и рецептур
  "содержит_контролируемые": false,                   # КВЭК: содержатся ли сведения из списков
  "требуется_лицензия": false,                        # КВЭК: нужна ли лицензия ФСТЭК / разрешение Комиссии
  "проверенные_пункты": [                             # КВЭК, обязательно, если что-то близко: пункты списков, под
    "п. 3.5.3 «б» ПП РФ № 1299 (…порог или суть…): в материале …; не подпадает, так как …"
  ],                                                  #   которые материал можно подтянуть, и почему он под них не подпадает
  "подразделение": "ЛМСТ",                            # ОП
  "издательство": "…", "город": "…", "страна": "…",   # ОП
  "проверенные_пункты_ОП": [                          # ОП, обязательно, если что-то близко: пункты Указа № 1203,
    "п. 63 Перечня, утв. Указом № 1203 (…суть…): в материале …; не подпадает, так как …"
  ],                                                  #   приказа ФСБ № 547 и т. п. и почему материал под них не подпадает
  "год": "2026"
}
"""
import copy
import json
import re
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

HERE = Path(__file__).resolve().parent
T_KVEK = HERE / 'АКТ КВЭК (шаблон).docx'
T_OP = HERE / 'ОП (шаблон 2026).docx'
OP_SAMPLE = HERE / 'образцы_2026' / 'Zaklyuchenie_po_OP_shablon_2026_s_25_05_2026.docx'
UND = re.compile(r'_{3,}')
FIO = 'И.О. Фамилия'


def paragraphs(d):
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def replace_span(p, start, end, new):
    """Заменить символы [start, end) текста абзаца на new, не трогая оформление вне этого участка."""
    pos = 0
    first = None
    for r in p.runs:
        t = r.text
        a, b = pos, pos + len(t)
        pos = b
        if b <= start or a >= end:
            continue
        lo, hi = max(start, a) - a, min(end, b) - a
        if first is None:
            r.text = t[:lo] + new + t[hi:]
            first = r
        else:
            r.text = t[:lo] + t[hi:]
    if first is None:
        # пустой участок в конце абзаца: дописать в последний прогон
        if start == end == len(p.text) and p.runs:
            p.runs[-1].text += new
            return
        raise ValueError(f'участок {start}:{end} не найден в «{p.text[:60]}»')


def fill_after(p, label, value):
    """Первую цепочку подчёркиваний после метки label заменить значением (бланк КВЭК)."""
    i = p.text.find(label)
    m = UND.search(p.text, i + len(label))
    tail = '' if p.text[m.end():m.end() + 1] in (',', '.', ';', ':') else ' '
    replace_span(p, m.start(), m.end(), ' ' + value + tail)


def fill_plain(p, old, new):
    i = p.text.find(old)
    if i < 0:
        raise ValueError(f'нет «{old}» в «{p.text[:80]}»')
    replace_span(p, i, i + len(old), new)


def set_value(p, label, value, stop=r'\t*_{3,}|\s*\(если|$'):
    """Форма 2026: значение после метки — до подчёркиваний, подсказки «(если…» или конца абзаца."""
    t = p.text
    i = t.find(label)
    if i < 0:
        raise ValueError(f'нет метки «{label}» в «{t[:80]}»')
    a = i + len(label)
    b = a + re.search(stop, t[a:]).start()
    replace_span(p, a, b, (' ' + value) if value else ' ')


def find(ps, prefix):
    for p in ps:
        if re.sub(r'\s+', ' ', p.text).strip().startswith(prefix):
            return p
    raise ValueError(f'нет абзаца «{prefix}»')


def make_op_template():
    """Пустой шаблон ОП 2026: поля материала и состав комиссии очищены, подписи — «И.О. Фамилия»."""
    d = docx.Document(OP_SAMPLE)
    ps = [p for p in paragraphs(d) if p.text.strip()]
    for label in ('Название работы:', 'Автор(ы):', 'Структурное подразделение:', 'Вид материала:',
                  'Организация / Издательство:'):
        set_value(find(ps, label), label, '')
    lab = 'Название конференции/ журнала/ монографии (др.) (выбрать нужное):'
    set_value(find(ps, lab), lab, '')
    jp = ps[ps.index(find(ps, lab)) + 1]                     # строка с названием журнала
    set_value(jp, '', '', stop=r'_{3,}')
    g = find(ps, 'Город:')
    set_value(g, 'Город:', '', stop=r'\.')
    set_value(g, 'Страна:', '', stop=r'\.')
    k = find(ps, 'Экспертная комиссия в составе')
    set_value(k, 'Экспертная комиссия в составе:', '______________________________________________________')
    for who in ('А.С. Перин', 'Т.Р. Газизов', 'М.Е. Комнатнов'):
        for p in ps:
            if who in p.text:
                fill_plain(p, who, FIO)
    d.save(T_OP)
    return T_OP


def add_checked_items(p, items, head='Материал проверен на соответствие близким к нему по предмету пунктам списков.'):
    """После абзаца-вывода (КВЭК: «подтверждаю, что … (не) содержится …»; ОП: «Сведения, содержащиеся в рассматриваемых
    материалах, не подпадают …») вставить пункты, под которые материал можно было бы подтянуть, с обоснованием, почему
    он под них не подпадает (требование автора 02.10.2026). Оформление — копия pPr абзаца и свойств первого прогона."""
    if not items:
        return
    anchor = p._p
    for text in [head] + [f'{i}) {s}' for i, s in enumerate(items, 1)]:
        new = copy.deepcopy(p._p)
        for child in list(new):
            if child.tag != qn('w:pPr'):
                new.remove(child)
        r = copy.deepcopy(p.runs[0]._r)
        for child in list(r):
            if child.tag != qn('w:rPr'):
                r.remove(child)
        new.append(r)
        anchor.addnext(new)
        anchor = new
        Paragraph(new, p._parent).runs[0].text = text


def fill_kvek(f, out):
    d = docx.Document(T_KVEK)
    ps = list(paragraphs(d))
    name = f['название'] + (f' ({f["перевод_названия"]})' if f.get('перевод_названия') else '')
    choice = '(выбрать необходимое, ненужное удалить)'
    dest = f.get('издание_для') or f['издание']
    for i, p in enumerate(ps):
        t = p.text
        if re.search(r'«_+»\s*_+\s*20\d\d', t):
            fill_plain(p, re.search(r'20\d\d', t).group(0), f.get('год', '2026'))
        elif t.startswith('рассмотрев материал'):
            fill_after(p, 'рассмотрев материал', f['вид'])
        elif t.startswith('наименование материала'):
            fill_after(p, 'наименование материала', name)
        elif t.startswith('автор(ы)'):
            fill_after(p, 'автор(ы)', f['авторы'])
        elif t.startswith('предназначенный для'):
            fill_after(p, 'предназначенный для', dest)
        elif 'краткое содержание' in t:
            fill_after(p, 'краткое содержание', f['краткое_содержание'])
            for q in ps[i + 1:i + 3]:
                if q.text.strip() and not re.sub(r'[_\s.]', '', q.text):
                    while UND.search(q.text):
                        m = UND.search(q.text)
                        replace_span(q, m.start(), m.end(), '')
        elif 'не содержится/содержится' in t:
            fill_plain(p, 'не содержится/содержится ' + choice,
                       'содержится' if f['содержит_контролируемые'] else 'не содержится')
            add_checked_items(p, f.get('проверенные_пункты') or [])
        elif t.startswith('Заключение:'):
            fill_plain(p, '(название конференции, журнала и т.д.)', dest)
            fill_plain(p, 'не требуется/требуется ' + choice,
                       'требуется' if f['требуется_лицензия'] else 'не требуется')
    path = Path(out) / f'КВЭК_{f["файл"]}.docx'
    d.save(path)
    return path


def fill_op(f, out):
    d = docx.Document(T_OP)
    ps = [p for p in paragraphs(d) if p.text.strip()]
    name = f['название'] + (f' ({f["перевод_названия"]})' if f.get('перевод_названия') else '')
    kind, _, journal = f['издание'].partition(' ')         # «журнал «Доклады ТУСУР»» -> «журнал», «Доклады ТУСУР»
    set_value(find(ps, 'Название работы:'), 'Название работы:', name)
    set_value(find(ps, 'Автор(ы):'), 'Автор(ы):', f['авторы'])
    set_value(find(ps, 'Структурное подразделение:'), 'Структурное подразделение:', f.get('подразделение', ''))
    lab = 'Название конференции/ журнала/ монографии (др.) (выбрать нужное):'
    set_value(find(ps, lab), lab, kind)
    jname = journal.strip('«»') if journal.count('«') == 1 else journal     # в образце журнал без кавычек
    set_value(ps[ps.index(find(ps, lab)) + 1], '', jname + '    ', stop=r'_{3,}')
    set_value(find(ps, 'Вид материала:'), 'Вид материала:', f.get('вид_ОП') or f['вид'])
    set_value(find(ps, 'Организация / Издательство:'), 'Организация / Издательство:', f.get('издательство', ''))
    g = find(ps, 'Город:')
    set_value(g, 'Город:', f.get('город', ''), stop=r'\.')
    set_value(g, 'Страна:', f.get('страна', ''), stop=r'\.')
    add_checked_items(find(ps, 'Сведения, содержащиеся в рассматриваемых материалах'), f.get('проверенные_пункты_ОП') or [],
                      head='Материал проверен на соответствие близким к нему по предмету пунктам перечней сведений, '
                           'составляющих государственную тайну, и сведений в военной и военно-технической области.')
    path = Path(out) / f'ОП_{f["файл"]}.docx'
    d.save(path)
    return path


def selfcheck():
    import tempfile
    f = dict(файл='тест', вид='статья', вид_ОП='статья в журнал', название='Методика X', авторы='Иванов А.Б.',
             издание='журнал «Y»', издание_для='журнала «Y»', краткое_содержание='Кратко.',
             содержит_контролируемые=False, требуется_лицензия=False, подразделение='ЛМСТ',
             проверенные_пункты=['п. 9.9 ПП РФ № 1299 (порог 1 ГГц): в материале 2 ГГц; не подпадает, так как …'],
             проверенные_пункты_ОП=['п. 99 Перечня, утв. Указом № 1203 (суть): в материале …; не подпадает, так как …'],
             издательство='ТУСУР', город='Томск', страна='Россия', год='2026')
    flat = lambda d: re.sub(r'\s+', ' ', '\n'.join(p.text for p in paragraphs(d)))
    with tempfile.TemporaryDirectory() as td:
        kt = flat(docx.Document(fill_kvek(f, td)))
        for s in ('рассмотрев материал статья', 'наименование материала Методика X', 'автор(ы) Иванов А.Б.',
                  'предназначенный для журнала «Y»', 'краткое содержание Кратко.', 'не содержится сведений',
                  'материалов для журнала «Y» не требуется оформление', '2026 г.', FIO):
            assert s in kt, s
        assert 'не содержится/содержится' not in kt and 'не требуется/требуется' not in kt
        i_ok, i_head, i_item, i_end = (kt.index(s) for s in ('не содержится сведений', 'близким к нему по предмету',
                                                             '1) п. 9.9 ПП РФ № 1299', 'Заключение:'))
        assert i_ok < i_head < i_item < i_end, 'проверенные пункты — после «подтверждаю» и до «Заключения»'
        blank = flat(docx.Document(T_OP))
        for gone in ('Шумы AlGaN', 'Моховиков', 'Перин', 'Газизов', 'Комнатнов', 'ЛМСТ', 'Москва', 'обзорная'):
            assert gone not in blank, f'в пустом шаблоне ОП осталось «{gone}»'
        ot = flat(docx.Document(fill_op(f, td)))
        for s in ('Название работы: Методика X', 'Автор(ы): Иванов А.Б.', 'Структурное подразделение: ЛМСТ',
                  '(выбрать нужное): журнал', ' Y ', 'Вид материала: статья в журнал',
                  'Организация / Издательство: ТУСУР', 'Город: Томск. Страна: Россия.',
                  'ред. от 24.06.2025', 'Председатель экспертной комиссии', FIO):
            assert s in ot, s
        assert ot.count(FIO) == 3, 'три подписи комиссии — пустые'
        j_ok, j_head, j_item, j_sig = (ot.index(s) for s in ('Сведения, содержащиеся в рассматриваемых материалах',
                                                             'составляющих государственную тайну, и сведений в военной',
                                                             '1) п. 99 Перечня', 'Председатель экспертной комиссии'))
        assert j_ok < j_head < j_item < j_sig, 'пункты ОП — после вывода комиссии и до подписей'
    print('selfcheck ок')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    if sys.argv[1] == '--selfcheck':
        selfcheck()
    elif sys.argv[1] == '--make-op-template':
        print(make_op_template())
    else:
        f = json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'))
        print(fill_kvek(f, sys.argv[2]))
        print(fill_op(f, sys.argv[2]))
